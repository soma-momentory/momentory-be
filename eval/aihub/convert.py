#!/usr/bin/env python3
"""
AI Hub 018.감성대화 → 모멘토리 정본(v2) 택소노미로 라벨 변환.

`mapping.json` 의 소분류 60종 → 대표 감정 12종 매핑을 적용하고, **커버리지를 먼저 보고한다.**
변환 자체보다 "무엇이 안 덮이는가"가 이 스크립트의 산출물이다.

    python3 convert.py --dir "…/018.감성대화" --report
    python3 convert.py --dir "…/018.감성대화" --out ../data/aihub-v2.jsonl \\
      --confidence direct --split validation

⚠ **이 결과는 v2 골드가 아니다.** AI Hub 는 대화 1건에 감정 **1개**만 달려 있고
강도·사건·근거 발화·명시적 표현 여부가 전부 없다. 정본 라벨링 단위(사건 × 감정)와 맞지 않는다.
계획 §6.1 대로 **V2 워밍업 자산**이고, 모델 비교(§9)의 정답지로 쓰지 않는다.

⚠ **AI Hub 이용 약관을 확인한 뒤 쓴다.** 2차 저작물·연구 외 활용에 제약이 있을 수 있다.
"""

import argparse
import json
import os
import sys
from collections import Counter, defaultdict

KO = {"joy": "기쁨", "calm": "평온", "proud": "뿌듯함", "affection": "애정", "sad": "슬픔",
      "anxious": "불안", "angry": "분노", "disappointed": "실망", "ashamed": "부끄러움",
      "confused": "혼란", "sorry": "미안함", "exhausted": "지침"}
ORDER = ["joy", "calm", "proud", "affection", "sad", "anxious", "angry", "disappointed",
         "ashamed", "confused", "sorry", "exhausted"]

SPLITS = {
    "training": ("Training_221115_add", "Training"),
    "validation": ("Validation_221115_add", "Validation"),
}


def load_mapping(path):
    d = json.load(open(path, encoding="utf-8"))
    return d["map"], d.get("out_reason", {}), d.get("unreachable", {})


def read_split(root, split):
    """
    xlsx(한글 라벨)와 json(대화 원문)을 **첫 사람 발화(HS01)로 조인**한다.

    행 순서로 이으면 안 된다 — 실측(2026-09-27): Training 은 xlsx 51,630행 / json 51,628건,
    Validation 은 6,641 / 6,640 으로 **xlsx 에만 있는 행이 섞여 있고 그 뒤로 순서가 어긋난다**
    (Validation 기준 index 490 부터). 순서를 믿으면 그 지점 이후 라벨이 전부 밀린다.

    HS01 이 겹치는 경우(Training 26종)는 **겹치는 것끼리 등장 순서로** 짝짓는다.
    """
    import openpyxl
    folder, name = SPLITS[split]
    xlsx = os.path.join(root, folder, "원천데이터", f"감성대화말뭉치(최종데이터)_{name}.xlsx")
    js = os.path.join(root, folder, "라벨링데이터", f"감성대화말뭉치(최종데이터)_{name}.json")

    talks = json.load(open(js, encoding="utf-8"))
    pool = defaultdict(list)                     # HS01 → [talk, ...] (등장 순서)
    for t in talks:
        c = (t.get("talk") or {}).get("content") or {}
        pool[c.get("HS01", "").strip()].append(t)

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    hi = {h: i for i, h in enumerate(hdr) if h}
    for k in ("감정_대분류", "감정_소분류", "사람문장1", "연령", "상황키워드"):
        if k not in hi:
            sys.exit(f"xlsx 에 '{k}' 열이 없습니다: {xlsx}")

    out, orphan = [], 0
    for r in it:
        if r[hi["감정_대분류"]] is None:
            continue
        first = (r[hi["사람문장1"]] or "").strip()
        bucket = pool.get(first)
        if not bucket:
            orphan += 1                          # json 에 짝이 없는 xlsx 행 — 버린다
            continue
        t = bucket.pop(0)
        c = (t.get("talk") or {}).get("content") or {}
        turns = [(k, c.get(k, "").strip())
                 for k in ("HS01", "SS01", "HS02", "SS02", "HS03", "SS03")]
        out.append({
            "major": str(r[hi["감정_대분류"]]).strip(),
            "minor": str(r[hi["감정_소분류"]]).strip(),
            "age": str(r[hi["연령"]] or "").strip(),
            "topic": str(r[hi["상황키워드"]] or "").strip(),
            "id": ((t.get("talk") or {}).get("id") or {}).get("talk-id"),
            "turns": [(k, v) for k, v in turns if v],
        })
    wb.close()

    leftover = sum(len(v) for v in pool.values())
    if orphan or leftover:
        print(f"  [{split}] 조인: {len(out)}건 · xlsx 짝 없음 {orphan} · json 짝 없음 {leftover}",
              file=sys.stderr)
    return out


def main():
    ap = argparse.ArgumentParser(description="AI Hub 018 → 정본 v2 라벨 변환")
    ap.add_argument("--dir", required=True, help="018.감성대화 폴더 경로")
    ap.add_argument("--mapping", default=os.path.join(os.path.dirname(__file__), "mapping.json"))
    ap.add_argument("--split", choices=["training", "validation", "both"], default="both")
    ap.add_argument("--out", help="변환 결과를 쓸 JSONL")
    ap.add_argument("--confidence", choices=["direct", "direct+approx", "all"],
                    default="direct+approx", help="어디까지 내보낼지 (기본 direct+approx)")
    ap.add_argument("--report", action="store_true", help="커버리지 표만 낸다")
    args = ap.parse_args()

    mp, out_reason, unreachable = load_mapping(args.mapping)
    splits = ["training", "validation"] if args.split == "both" else [args.split]

    rows = []
    for s in splits:
        rows += [{**r, "split": s} for r in read_split(args.dir, s)]

    per_emotion = defaultdict(Counter)   # v2 감정 -> confidence -> n
    dropped = Counter()
    unmapped = Counter()
    kept = []

    allow = {"direct"} if args.confidence == "direct" else \
            {"direct", "approx"} if args.confidence == "direct+approx" else \
            {"direct", "approx", "out"}

    for r in rows:
        key = f"{r['major']}|{r['minor']}"
        hit = mp.get(key)
        if hit is None:
            unmapped[key] += 1
            continue
        label, conf = hit
        if label is None:
            dropped[key] += 1
            if "out" not in allow:
                continue
        else:
            per_emotion[label][conf] += 1
        if conf in allow:
            kept.append({
                "sourceId": r["id"], "split": r["split"],
                "taxonomy": "v2", "source": "aihub",
                "normalized": label, "confidence": conf,
                "aihub": {"major": r["major"], "minor": r["minor"],
                          "age": r["age"], "topic": r["topic"]},
                "turns": [{"speaker": "user" if k.startswith("HS") else "assistant", "text": v}
                          for k, v in r["turns"]],
            })

    total = len(rows)
    print(f"원본 {total}건 ({', '.join(splits)})")
    if unmapped:
        print(f"\n⚠ 매핑에 없는 소분류 {len(unmapped)}종 — mapping.json 을 고쳐야 합니다:")
        for k, n in unmapped.most_common():
            print(f"    {k}  {n}건")

    print(f"\n── 정본 12종 커버리지 ──")
    print(f"  {'감정':<6} {'direct':>8} {'approx':>8} {'합계':>8}")
    for e in ORDER:
        d, a = per_emotion[e]["direct"], per_emotion[e]["approx"]
        mark = "   ← 없음" if d + a == 0 else ("   ← direct 없음" if d == 0 else "")
        print(f"  {KO[e]:<6} {d:>8} {a:>8} {d + a:>8}{mark}")
    mapped = sum(sum(c.values()) for c in per_emotion.values())
    out_n = sum(dropped.values())
    print(f"\n  매핑됨 {mapped}건 ({mapped / total:.1%}) · 체계 밖 {out_n}건 ({out_n / total:.1%})")

    if unreachable:
        print("\n── 이 데이터로 못 채우는 감정 ──")
        for k, why in unreachable.items():
            print(f"  {KO.get(k, k)}: {why}")

    print("\n── 체계 밖으로 뺀 소분류 ──")
    for k, n in dropped.most_common():
        print(f"  {k:<22} {n:>6}건  {out_reason.get(k, '')}")

    if args.out and not args.report:
        with open(args.out, "w", encoding="utf-8") as f:
            for r in kept:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        print(f"\n{len(kept)}건 → {args.out}  (confidence={args.confidence})")
        print("⚠ v2 골드가 아니다 — 감정 1개뿐이고 강도·사건·근거가 없다. V2 워밍업 자산으로만 쓴다.")


if __name__ == "__main__":
    main()
